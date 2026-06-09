from datetime import datetime
import os
import re
import math
import openpyxl
from elasticsearch import Elasticsearch

class LogService:
    def __init__(self):
        self.es = Elasticsearch(os.getenv("ELASTICSEARCH_HOST"))
        self.index = "log-files"
        self.ensure_index()

    def ensure_index(self):
        exists = self.es.indices.exists(index=self.index)
        if exists:
            # Check if mapping contains filename_event, if not delete and recreate
            try:
                mapping = self.es.indices.get_mapping(index=self.index)
                props = mapping.get(self.index, {}).get("mappings", {}).get("properties", {})
                if "filename_event" not in props:
                    self.es.indices.delete(index=self.index)
                    exists = False
            except Exception:
                # Fallback to delete and recreate if any issue reading mapping
                self.es.indices.delete(index=self.index, ignore_unavailable=True)
                exists = False

        if not exists:
            self.es.indices.create(
                index=self.index,
                mappings={
                    "properties": {
                        "filename_event": {"type": "keyword"},
                        "filename_log":   {"type": "keyword"},
                        "app_name":       {"type": "keyword"},
                        "pair_id":        {"type": "keyword"},
                        "event_content":  {
                            "type": "text",
                            "analyzer": "standard",
                        },
                        "log_content":    {
                            "type": "text",
                            "analyzer": "log_analyzer",
                            "term_vector": "with_positions_offsets",
                            "fields": {
                                "standard": {
                                    "type": "text",
                                    "analyzer": "standard"
                                }
                            }
                        },
                        "line_count": {"type": "integer"},
                        "file_size":  {"type": "long"},
                        "uploaded_at": {"type": "date"},
                    }
                },
                settings={
                    "analysis": {
                        "analyzer": {
                            "log_analyzer": {
                                "type": "custom",
                                "tokenizer": "standard",
                                "filter": ["lowercase", "log_stop_words"],
                                "char_filter": ["log_normalizer"],
                            }
                        },
                        "char_filter": {
                            "log_normalizer": {
                                "type": "pattern_replace",
                                "pattern": r"(\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}|\[\d+\]|\d{2}:\d{2}:\d{2}|port \d+)",
                                "replacement": " ",
                            }
                        },
                        "filter": {
                            "log_stop_words": {
                                "type": "stop",
                                "stopwords": ["the", "a", "an", "is", "by", "for", "from", "to", "at", "in", "of"],
                            }
                        },
                    }
                },
            )

    def _count_lines(self, content):
        return sum(1 for line in content.splitlines() if line.strip())

    def upload_pair(self, event_file, log_file, app_name=None, pair_id=None):
        event_content = event_file.read().decode("utf-8", errors="ignore")
        log_content = log_file.read().decode("utf-8", errors="ignore")

        # Deduce metadata
        if not pair_id:
            for name in [log_file.filename, event_file.filename]:
                if name:
                    match = re.search(r'(?:log|issue)_([a-zA-Z0-9_-]+)\.', name)
                    if match:
                        pair_id = match.group(1)
                        break
        if not app_name:
            for name in [log_file.filename, event_file.filename]:
                if name:
                    parts = name.replace('\\', '/').split('/')
                    if len(parts) > 1:
                        app_name = parts[0]
                        break
            if not app_name:
                app_name = 'unknown'
        if not pair_id:
            pair_id = 'unknown'

        filename_event = event_file.filename or f"issue_{pair_id}.md"
        filename_log = log_file.filename or f"log_{pair_id}.log"
        
        line_count = self._count_lines(log_content)
        file_size = len(log_content.encode("utf-8")) + len(event_content.encode("utf-8"))

        doc_body = {
            "filename_event": filename_event,
            "filename_log": filename_log,
            "app_name": app_name,
            "pair_id": pair_id,
            "event_content": event_content,
            "log_content": log_content,
            "line_count": line_count,
            "file_size": file_size,
            "uploaded_at": datetime.utcnow().isoformat() + "Z",
        }

        # Use filename_log as document id to prevent duplicates
        self.es.index(index=self.index, id=filename_log, document=doc_body)
        self.es.indices.refresh(index=self.index)

        return {
            "filename_event": filename_event,
            "filename_log": filename_log,
            "app_name": app_name,
            "pair_id": pair_id,
            "action": "indexed",
            "lineCount": line_count
        }

    def list_files(self):
        result = self.es.search(
            index=self.index,
            size=100,
            query={"match_all": {}},
            _source=["filename_event", "filename_log", "app_name", "pair_id", "line_count", "file_size", "uploaded_at", "event_content", "log_content"],
            sort=[{"uploaded_at": {"order": "desc"}}],
        )
        return [
            {
                "id": hit["_id"],
                **hit["_source"],
            }
            for hit in result["hits"]["hits"]
        ]

    def delete_file(self, file_id):
        self.es.delete(index=self.index, id=file_id)
        self.es.indices.refresh(index=self.index)

    def search(self, query_text, search_type='log', size=10):
        target_field = "log_content" if search_type == "log" else "event_content"
        
        result = self.es.search(
            index=self.index,
            size=size,
            query={
                "match": {
                    target_field: {
                        "query": query_text,
                        "operator": "or",
                        "fuzziness": "AUTO",
                    }
                }
            },
            highlight={
                "fields": {
                    target_field: {
                        "fragment_size": 200,
                        "number_of_fragments": 5,
                        "pre_tags": ["<<<"],
                        "post_tags": [">>>"],
                    }
                }
            },
            _source=["filename_event", "filename_log", "app_name", "pair_id", "line_count", "file_size", "event_content", "log_content"],
        )

        hits = []
        for hit in result["hits"]["hits"]:
            highlights = [
                {
                    "text": fragment.replace("<<<", "").replace(">>>", "").strip(),
                    "marked": fragment,
                }
                for fragment in hit.get("highlight", {}).get(target_field, [])
            ]

            hits.append({
                "id": hit["_id"],
                "filename_event": hit["_source"].get("filename_event"),
                "filename_log": hit["_source"].get("filename_log"),
                "app_name": hit["_source"].get("app_name"),
                "pair_id": hit["_source"].get("pair_id"),
                "line_count": hit["_source"].get("line_count"),
                "file_size": hit["_source"].get("file_size"),
                "event_content": hit["_source"].get("event_content"),
                "log_content": hit["_source"].get("log_content"),
                "score": hit.get("_score"),
                "highlights": highlights,
            })

        return {
            "total": result["hits"]["total"]["value"],
            "max_score": result["hits"]["max_score"],
            "hits": hits,
        }

    def import_pairs(self):
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "pares_logs_issues", "out"))
        if not os.path.exists(base_dir):
            base_dir = os.path.abspath("pares_logs_issues/out")
            
        if not os.path.exists(base_dir):
            raise FileNotFoundError(f"Directory not found: {base_dir}")

        imported_count = 0
        
        for root, dirs, files in os.walk(base_dir):
            for file in files:
                if file.startswith("issue_") and file.endswith(".md"):
                    pair_id = file[len("issue_"):-len(".md")]
                    log_file_name = f"log_{pair_id}.log"
                    if log_file_name in files:
                        event_path = os.path.join(root, file)
                        log_path = os.path.join(root, log_file_name)
                        
                        rel_event_path = os.path.relpath(event_path, base_dir)
                        rel_log_path = os.path.relpath(log_path, base_dir)
                        
                        app_name = os.path.basename(root)
                        
                        with open(event_path, "r", encoding="utf-8", errors="ignore") as f:
                            event_content = f.read()
                        with open(log_path, "r", encoding="utf-8", errors="ignore") as f:
                            log_content = f.read()
                            
                        line_count = self._count_lines(log_content)
                        file_size = len(log_content.encode("utf-8")) + len(event_content.encode("utf-8"))
                        
                        doc_body = {
                            "filename_event": rel_event_path,
                            "filename_log": rel_log_path,
                            "app_name": app_name,
                            "pair_id": pair_id,
                            "event_content": event_content,
                            "log_content": log_content,
                            "line_count": line_count,
                            "file_size": file_size,
                            "uploaded_at": datetime.utcnow().isoformat() + "Z",
                        }
                        
                        self.es.index(index=self.index, id=rel_log_path, document=doc_body)
                        imported_count += 1
                        
        self.es.indices.refresh(index=self.index)
        return imported_count

    def calculate_metrics(self):
        excel_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "relevancia_esperada_logs_116_118_285.xlsx"))
        if not os.path.exists(excel_path):
            excel_path = os.path.abspath("relevancia_esperada_logs_116_118_285.xlsx")
            
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Excel file not found at {excel_path}")
            
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        
        queries_data = {}
        for r in range(2, ws.max_row + 1):
            query_text = ws.cell(row=r, column=2).value
            rank = ws.cell(row=r, column=3).value
            expected_log = ws.cell(row=r, column=4).value
            
            if not query_text or not expected_log:
                continue
                
            query_text = query_text.strip()
            expected_log = expected_log.replace("\n", "").replace("\r", "").strip().replace("\\", "/")
            rank = int(rank)
            
            if query_text not in queries_data:
                queries_data[query_text] = {}
            queries_data[query_text][rank] = expected_log

        strategies = {
            "Log Search (Standard)": {
                "match": {
                    "log_content.standard": {
                        "operator": "or",
                        "fuzziness": "AUTO"
                    }
                }
            },
            "Log Search (Custom Analyzer)": {
                "match": {
                    "log_content": {
                        "operator": "or",
                        "fuzziness": "AUTO"
                    }
                }
            },
            "Event Search (Standard)": {
                "match": {
                    "event_content": {
                        "operator": "or",
                        "fuzziness": "AUTO"
                    }
                }
            },
            "Hybrid Search (Multi-match on both fields)": {
                "multi_match": {
                    "fields": ["log_content", "event_content"],
                    "type": "best_fields",
                    "operator": "or",
                    "fuzziness": "AUTO"
                }
            }
        }
        
        idcg_5 = 0.0
        for i in range(1, 6):
            rel = 6 - i
            idcg_5 += (2**rel - 1) / math.log2(i + 1)
            
        results = {}
        
        for strategy_name, query_template in strategies.items():
            rr_sum = 0.0
            ndcg_sum = 0.0
            query_details = []
            
            for query_text, expected_ranks in queries_data.items():
                target_log = expected_ranks.get(1)
                
                if "match" in query_template:
                    field = list(query_template["match"].keys())[0]
                    body = {
                        "query": {
                            "match": {
                                field: {
                                    "query": query_text,
                                    **query_template["match"][field]
                                }
                            }
                        }
                    }
                else:
                    body = {
                        "query": {
                            "multi_match": {
                                "query": query_text,
                                **query_template["multi_match"]
                            }
                        }
                    }
                
                res = self.es.search(index=self.index, size=50, query=body["query"], _source=["filename_log"])
                hits = res["hits"]["hits"]
                
                retrieved_logs = []
                for hit in hits:
                    fn = hit["_source"].get("filename_log", "")
                    fn = fn.replace("\\", "/").strip()
                    retrieved_logs.append(fn)
                
                rr = 0.0
                if target_log:
                    target_log_clean = target_log.replace("\\", "/").strip()
                    if target_log_clean in retrieved_logs:
                        rank_found = retrieved_logs.index(target_log_clean) + 1
                        rr = 1.0 / rank_found
                
                dcg_5 = 0.0
                for i in range(min(5, len(retrieved_logs))):
                    ret_fn = retrieved_logs[i]
                    matched_rank = None
                    for exp_rank, exp_fn in expected_ranks.items():
                        if exp_fn.replace("\\", "/").strip() == ret_fn:
                            matched_rank = exp_rank
                            break
                    
                    rel = 0
                    if matched_rank is not None:
                        rel = 6 - matched_rank
                        
                    dcg_5 += (2**rel - 1) / math.log2(i + 2)
                
                ndcg_5 = dcg_5 / idcg_5 if idcg_5 > 0 else 0.0
                
                rr_sum += rr
                ndcg_sum += ndcg_5
                
                query_details.append({
                    "query": query_text,
                    "target_log": target_log,
                    "rr": rr,
                    "ndcg_5": ndcg_5,
                    "retrieved_top_5": retrieved_logs[:5]
                })
                
            num_queries = len(queries_data)
            mrr = rr_sum / num_queries if num_queries > 0 else 0.0
            mndcg_5 = ndcg_sum / num_queries if num_queries > 0 else 0.0
            
            results[strategy_name] = {
                "mrr": mrr,
                "ndcg_5": mndcg_5,
                "queries": query_details
            }
            
        return results

    def health(self):
        health = self.es.cluster.health()
        stats = self.es.count(index=self.index)
        return {
            "status": health.get("status"),
            "documents": stats.get("count"),
        }
