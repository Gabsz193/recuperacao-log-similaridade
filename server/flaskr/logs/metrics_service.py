import os
import openpyxl
import ir_measures
from ir_measures import nDCG, RR

class MetricsService:
    @staticmethod
    def calculate_metrics(log_service):
        excel_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "julgamento_relevancia_40_queries.xlsx"))
        if not os.path.exists(excel_path):
            excel_path = os.path.abspath("julgamento_relevancia_40_queries.xlsx")
            
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Excel file not found at {excel_path}")
            
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        
        queries_data = {}
        for r in range(2, ws.max_row + 1):
            log_alvo = ws.cell(row=r, column=1).value
            query_text = ws.cell(row=r, column=2).value
            rank = ws.cell(row=r, column=3).value
            expected_log = ws.cell(row=r, column=4).value
            
            if not query_text or not expected_log:
                continue
                
            log_alvo = log_alvo.strip() if log_alvo else ""
            query_text = query_text.strip()
            expected_log = expected_log.replace("\n", "").replace("\r", "").strip().replace("\\", "/")
            rank = int(rank)
            
            key = (log_alvo, query_text)
            if key not in queries_data:
                queries_data[key] = {}
            queries_data[key][rank] = expected_log

        strategies = {
            "Log Search (Standard)": log_service.search_log_standard,
            "Log Search (Custom Analyzer)": log_service.search_log_custom,
            "Event Search (Standard)": log_service.search_event_standard,
            "Hybrid Search (Multi-match on both fields)": log_service.search_hybrid,
        }
        
        # Build qrels and query_map once for ir-measures
        qrels = {}
        query_map = {}
        
        for idx, ((log_alvo, query_text), expected_ranks) in enumerate(queries_data.items()):
            q_id = f"q_{idx}"
            query_map[q_id] = {
                "log_alvo": log_alvo,
                "query_text": query_text,
                "expected_ranks": expected_ranks,
                "target_log": expected_ranks.get(1)
            }
            
            qrels[q_id] = {}
            for rank, expected_log in expected_ranks.items():
                qrels[q_id][expected_log] = 6 - rank
                
        results = {}
        
        for strategy_name, search_func in strategies.items():
            run = {}
            retrieved_logs_map = {}
            
            for q_id, q_info in query_map.items():
                query_text = q_info["query_text"]
                res = search_func(query_text, size=50)
                
                run[q_id] = {}
                retrieved_logs = []
                for idx_hit, hit in enumerate(res["hits"]):
                    fn = hit.get("filename_log", "")
                    fn = fn.replace("\\", "/").strip()
                    retrieved_logs.append(fn)
                    # Enforce strict ES ranking order by using artificial scoring
                    run[q_id][fn] = 1000.0 - idx_hit
                
                retrieved_logs_map[q_id] = retrieved_logs
                
            # Compute global aggregate metrics using ir-measures
            res_agg = ir_measures.calc_aggregate([nDCG@5, RR(rel=5)], qrels, run)
            mrr = res_agg.get(RR(rel=5), 0.0)
            mndcg_5 = res_agg.get(nDCG@5, 0.0)
            
            # Compute per-query metrics using ir-measures
            query_metrics = {}
            for metric in ir_measures.iter_calc([nDCG@5, RR(rel=5)], qrels, run):
                q_id = metric.query_id
                if q_id not in query_metrics:
                    query_metrics[q_id] = {"rr": 0.0, "ndcg_5": 0.0}
                if metric.measure == RR(rel=5):
                    query_metrics[q_id]["rr"] = metric.value
                elif metric.measure == nDCG@5:
                    query_metrics[q_id]["ndcg_5"] = metric.value
                    
            query_details = []
            for q_id, q_info in query_map.items():
                metrics_val = query_metrics.get(q_id, {"rr": 0.0, "ndcg_5": 0.0})
                query_details.append({
                    "query": q_info["query_text"],
                    "target_log": q_info["target_log"],
                    "rr": metrics_val["rr"],
                    "ndcg_5": metrics_val["ndcg_5"],
                    "retrieved_top_5": retrieved_logs_map.get(q_id, [])[:5]
                })
                
            results[strategy_name] = {
                "mrr": mrr,
                "ndcg_5": mndcg_5,
                "queries": query_details
            }
            
        return results
